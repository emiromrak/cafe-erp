from _paths import IMAGE_DIR, OUTPUT_PATH, TEMPLATE_PATH

import docx

doc = docx.Document(TEMPLATE_PATH)

# Extract texts:
# 1.A:
t1 = doc.tables[1]
text_1a = t1.rows[1].cells[0].text.strip()
print("--- 1.A ÖZET ---")
print(text_1a[:100], "...")

# 1.B: P[35], P[36]
print("\n--- 1.B AMACI ---")
print(doc.paragraphs[35].text[:100])
print(doc.paragraphs[36].text[:100])

# 1.C: P[37], P[38]
print("\n--- 1.C KONU VE KAPSAM ---")
print(doc.paragraphs[37].text[:100])
print(doc.paragraphs[38].text[:100])

# 1.D: P[39], P[40]
print("\n--- 1.D LİTERATÜR ÖZETİ ---")
print(doc.paragraphs[39].text[:100])
print(doc.paragraphs[40].text[:100])

# 2.A: P[42]
print("\n--- 2.A ÖZGÜN DEĞER ---")
print(doc.paragraphs[42].text[:100])

# 2.B: P[49], P[50], P[51]
print("\n--- 2.B YÖNTEM ---")
print(doc.paragraphs[49].text[:100])
print(doc.paragraphs[50].text[:100])
print(doc.paragraphs[51].text[:100])

# 2.C: P[65] (img), P[80] (Grafik-1), P[81]..P[84]
print("\n--- 2.C ÇALIŞMA TAKVİMİ ---")
print(doc.paragraphs[80].text)
print(doc.paragraphs[81].text)
print(doc.paragraphs[82].text)
print(doc.paragraphs[83].text)
print(doc.paragraphs[84].text)

# 2.D: t7.rows[1]
t7 = doc.tables[7]
print("\n--- 2.D KAYNAKÇA ---")
print(t7.rows[1].cells[0].text[:100])

# 3.A: t7.rows[3]
print("\n--- 3.A SUNUM ---")
print(t7.rows[3].cells[0].text[:100])

# 3.B: t7.rows[5]
print("\n--- 3.B RAPOR DÜZENİ ---")
print(t7.rows[5].cells[0].text[:100])
