from _paths import IMAGE_DIR, OUTPUT_PATH, TEMPLATE_PATH

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

# Source document
src_path = TEMPLATE_PATH
dst_path = OUTPUT_PATH

doc = docx.Document(src_path)

# 1. Inspect cover page transition:
# Check if P[33] needs a page break before P[34]
p33 = doc.paragraphs[33]
p34 = doc.paragraphs[34]

# Ensure page break after cover page:
# In python-docx, add page break run to p33 or set p34.paragraph_format.page_break_before = True
p34.paragraph_format.page_break_before = True

print("Page break set on P[34].")
