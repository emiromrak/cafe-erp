from _paths import IMAGE_DIR, OUTPUT_PATH, TEMPLATE_PATH

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os

src_path = TEMPLATE_PATH
dst_path = OUTPUT_PATH

flowchart_img = IMAGE_DIR / "yontem_akis_semasi.png"
gantt_img = IMAGE_DIR / "calisma_takvimi_gantt.png"

doc = docx.Document(src_path)

# Let's inspect paragraphs 0 to 34 and Table 0
body = doc.element.body

# Identify elements up to Table 0
elements = list(body)
table0_idx = -1
for idx, el in enumerate(elements):
    if el.tag.endswith('tbl'):
        table0_idx = idx
        break

print(f"Table 0 found at body element index {table0_idx}")

# Set page break on P[34] so Page 2 starts cleanly
p34 = doc.paragraphs[34]
p34.paragraph_format.page_break_before = True

# Now, we want to replace all elements after Table 0 with our single, beautifully formatted table!
# Let's collect elements to remove
elements_to_remove = elements[table0_idx + 1 :]
print(f"Elements to remove after Table 0: {len(elements_to_remove)}")
for el in elements_to_remove:
    # Do not remove sectPr at the very end
    if not el.tag.endswith('sectPr'):
        body.remove(el)

print("Removed messy tables and stray paragraphs.")

# Now, create the unified table
# It has width 10588 dxa, centered, with standard borders
unified_table = doc.add_table(rows=0, cols=1)
tblPr = unified_table._tbl.tblPr

# Set table properties: width 10588, jc center, borders
tblW = parse_xml(r'<w:tblW {} w:w="10588" w:type="dxa"/>'.format(nsdecls('w')))
tblJc = parse_xml(r'<w:jc {} w:val="center"/>'.format(nsdecls('w')))
tblBorders = parse_xml(r'''
    <w:tblBorders {}>
        <w:top w:val="single" w:color="auto" w:sz="4" w:space="0"/>
        <w:bottom w:val="single" w:color="auto" w:sz="4" w:space="0"/>
        <w:insideH w:val="single" w:color="auto" w:sz="4" w:space="0"/>
        <w:insideV w:val="single" w:color="auto" w:sz="4" w:space="0"/>
    </w:tblBorders>
'''.format(nsdecls('w')))

tblPr.append(tblW)
tblPr.append(tblJc)
tblPr.append(tblBorders)

def add_header_row(table, title_text):
    row = table.add_row()
    # Row props: trHeight 567 dxa, cantSplit
    trPr = row._tr.get_or_add_trPr()
    trHeight = parse_xml(r'<w:trHeight {} w:val="567"/>'.format(nsdecls('w')))
    cantSplit = parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w')))
    trPr.append(trHeight)
    trPr.append(cantSplit)

    cell = row.cells[0]
    # Shading: BFBFBF
    shd = parse_xml(r'<w:shd {} w:val="clear" w:color="auto" w:fill="BFBFBF"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)

    # Cell margins / padding: top/bottom 80 dxa, left/right 120 dxa
    tcMar = parse_xml(r'''
        <w:tcMar {}>
            <w:top w:w="80" w:type="dxa"/>
            <w:bottom w:w="80" w:type="dxa"/>
            <w:left w:w="120" w:type="dxa"/>
            <w:right w:w="120" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(tcMar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(title_text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    return row

def add_content_row(table):
    row = table.add_row()
    # Row props: no cantSplit, allows natural page breaking
    cell = row.cells[0]
    # Shading auto (white)
    shd = parse_xml(r'<w:shd {} w:val="clear" w:color="auto" w:fill="auto"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)

    # Cell margins: top/bottom 120 dxa, left/right 150 dxa
    tcMar = parse_xml(r'''
        <w:tcMar {}>
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="150" w:type="dxa"/>
            <w:right w:w="150" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(tcMar)
    return cell

def add_body_paragraph(cell, text, is_first=False, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = cell.paragraphs[0] if is_first else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = align

    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = False
    return p

# 1.A. ÖZET
add_header_row(unified_table, "1.A. / TASARIM PROJESİNİN ÖZETİ SUMMARY OF DESIGN PROJECT [PUAN / SCORE : 5]")
c1a = add_content_row(unified_table)
text_1a_p1 = "Yeme-içme sektöründe çalışan küçük ve orta ölçekli işletmeler ile butik kafelerde, mevcut satış noktası (POS) yazılımları genellikle sadece sipariş alma ve adisyon kapatma işlemleri için kullanılmaktadır. Ürünlerin reçetelerine bağlı hammadde takibinin, anlık stok yönetiminin ve geleceğe yönelik sipariş planlamasının el yordamıyla yapılması; işletmelerde gereksiz ürün israfına, stok tükenmesine ve maliyet artışına sebep olmaktadır (Türker vd., 2026; Chen vd., 2024). Bu projede, kafe işletmelerinin operasyonel süreçlerini düzenleyen ve yapay zeka destekli karar mekanizması içeren yönetim sistemi geliştirilmiştir."
text_1a_p2 = "Proje kapsamında; FastAPI tabanlı backend servisi, ilişkisel veri tabanı yapısı, sipariş verildiği anda ilgili hammaddeleri gramaj ve mililitre bazında stoktan düşen reçete motoru ve çoklu istemci desteği (kasa terminali, mobil garson arayüzü ve mutfak ekranı) kurulmuştur. Karar destek tarafında ise geçmiş satış kayıtları ile harici hava durumu verilerini birlikte değerlendiren LightGBM makine öğrenmesi modeli kullanılarak ürün ve kategori bazında satış tahminleri üretilmektedir (Wang ve Zhang, 2025). Modelin ürettiği tüketim hızları ile mevcut depo stokları kıyaslanarak, karar motoru ve Telegram botu üzerinden işletmeciye anlık tükenme uyarıları ve malzeme sipariş önerileri aktarılmaktadır. Sistem, kafe operasyonlarını tek bir merkezden yürütmeyi ve veri odaklı kararlarla fire oranlarını düşürmeyi amaçlamaktadır."
add_body_paragraph(c1a, text_1a_p1, is_first=True)
add_body_paragraph(c1a, text_1a_p2, is_first=False)

# 1.B. AMACI
add_header_row(unified_table, "1.B. TASARIM PROJESİNİN AMACI / aım of DESIGN PROJECT [PUAN / SCORE : 10]")
c1b = add_content_row(unified_table)
text_1b_p1 = "Bu projenin temel amacı; kafe ve restoran işletmelerinde sıkça karşılaşılan plansız hammadde tedariki, yüksek fire oranları ve manuel stok takibi sorunlarını modern yazılım ve makine öğrenmesi yöntemleriyle çözmektir. Bu ana amaç doğrultusunda, satışı yapılan her menü ürününün ve seçilen varyasyonların reçetelerine göre kahve çekirdeği, süt, şurup gibi hammaddelerin depodaki miktarından anlık ve otomatik olarak düşülmesi hedeflenmektedir."
text_1b_p2 = "Bunun yanı sıra geçmiş satış örüntülerini, haftanın günlerini, saat dilimlerini ve günlük hava durumu verilerini işleyerek sonraki günlerin ürün talebini yüksek doğrulukla tahmin etmek amaçlanmaktadır (Aksoy ve Çelik, 2024). Tahmin edilen talep ile mevcut stok verileri birleştirilerek hammaddelerin tükenme sürelerinin hesaplanması ve kritik seviyeye yaklaşan ürünler için işletmeciye Telegram üzerinden doğrudan uyarı ve satın alma önerisi gönderilmesi hedeflenmektedir. Ayrıca kasa terminali, garsonların kendi telefonlarından sipariş girebileceği mobil arayüz ve mutfak takip ekranının seknron çalıştırılmasıyla kafe içi sipariş trafiğinin hızlandırılması amaçlanmaktadır."
add_body_paragraph(c1b, text_1b_p1, is_first=True)
add_body_paragraph(c1b, text_1b_p2, is_first=False)

# 1.C. KONU VE KAPSAM
add_header_row(unified_table, "1.C. KONU ve KAPSAM / subject and SCOPE [PUAN / SCORE : 15]")
c1c = add_content_row(unified_table)
text_1c_p1 = "Projenin konusu; yeme-içme işletmelerine yönelik, satış noktası operasyonları, dinamik reçete-stok takip mekanizması ve makine öğrenmesi tabanlı talep tahminleme özelliklerini tek bir yapıda toplayan bir işletim sisteminin tasarlanması ve uygulanmasıdır."
text_1c_p2 = "Çalışma kapsamında FastAPI ile REST mimarisine uygun ve hızlı çalışan backend servisleri geliştirilmiştir. Veri katmanında masalar, siparişler, menü bileşenleri, ürün varyasyonları, reçeteler ve hammadde envanterini tutan ilişkisel bir veri tabanı şeması kurulmuştur. Kullanıcı tarafında ise kasa terminali, mobil garson sipariş ekranı ve mutfak takip ekranı tasarlanmıştır. Sistemin karar destek boyutunda Open-Meteo servisinden alınan hava durumu verileri ile satış geçmişi birleştirilerek LightGBM modeli eğitilmiş, stok tükenme süresini hesaplayan kural tabanlı uzman motor ve işletmeciye mesaj ileten Telegram botu sisteme entegre edilmiştir. Sistemin doğruluğunu ve yanıt sürelerini ölçmek amacıyla hazırlanan benzetim veri setiyle kapsamlı fonksiyonel testler gerçekleştirilmiştir. Resmi mali mühür, e-fatura entegrasyonu ve yasal muhasebe defter kayıtları gibi yasal süreçler ise projenin operasyonel yönetim ve karar destek odağını korumak amacıyla ilk aşamada kapsam dışı bırakılmıştır."
add_body_paragraph(c1c, text_1c_p1, is_first=True)
add_body_paragraph(c1c, text_1c_p2, is_first=False)

# 1.D. LİTERATÜR ÖZETİ
add_header_row(unified_table, "1.D. LİTERATÜR ÖZETİ / SUMMARY OF LITERATURE [PUAN / SCORE : 10]")
c1d = add_content_row(unified_table)
text_1d_p1 = "Restoran ve kafe işletmeciliğinde dijital dönüşüm, uzun yıllar boyunca yalnızca masa takibi ve adisyon fişi kesen geleneksel satış noktası cihazları üzerinden yürütülmüştür (Kaya ve Yıldız, 2023). Ancak bu sistemler genellikle operasyonel satış verilerini depolamakta, bu verileri hammadde düzeyinde envanter yönetimiyle veya geleceğe yönelik planlamayla ilişkilendirmemektedir. Son yıllarda yapılan çalışmalar, yeme-içme sektöründeki gıda israfının ve maliyet kayıplarının önemli bir kısmının hatalı talep tahminleri ve kontrolsüz stok yönetiminden kaynaklandığını ortaya koymaktadır (Türker vd., 2026)."
text_1d_p2 = "Zaman serisi ve makine öğrenmesi algoritmalarının perakende sektöründe talep tahmininde kullanımı literatürde geniş yer bulmaktadır. Özellikle ağaç tabanlı gradyan artırma yöntemleri ve LightGBM algoritması; tablosal verilerde hızlı eğitim süresi, düşük bellek kullanımı ve yüksek tahmin doğruluğu ile öne çıkmaktadır (Ke vd., 2017). Yeme-içme talebinin hava durumu, sıcaklık, yağış ve günün saatleri gibi dış etkenlerden doğrudan etkilendiği bilinmektedir (Chen vd., 2024). Literatürde dış değişkenleri modele dahil eden çalışmaların, yalnızca geçmiş satış rakamlarına dayanan geleneksel istatistiksel yöntemlere göre belirgin şekilde daha düşük hata oranları verdiği kanıtlanmıştır (Wang ve Zhang, 2025). Diğer taraftan üretilen tahminlerin sahada hızlıca uygulanabilmesi için kural tabanlı bildirim mekanizmalarının ve anlık mesajlaşma platformlarının küçük işletmelerde uygulanabilirlik açısından yüksek verim sağladığı belirtilmektedir (Aksoy ve Çelik, 2024)."
add_body_paragraph(c1d, text_1d_p1, is_first=True)
add_body_paragraph(c1d, text_1d_p2, is_first=False)

# 2.A. ÖZGÜN DEĞER
add_header_row(unified_table, "2.A. ÖZGÜN DEĞER / ORIGINALITY [PUAN / SCORE : 10]")
c2a = add_content_row(unified_table)
text_2a = "Projenin özgün değeri; küçük ölçekli kafe işletmelerine yönelik hafif bir mimaride, hammadde düzeyinde reçete takibi ile yapay zeka tabanlı talep tahminini tek bir sistemde birleştirmesidir. Mevcut ticari paketler stok takibini yalnızca paketli ürün bazında yaparken; bu sistemde her sipariş varyasyonunun harcadığı gramaj bazlı hammadde reçeteden anlık düşülmektedir. Ayrıca hava durumu verileriyle beslenen yerel yapay zeka modelinin ürettiği satış öngörüleri, karmaşık grafikler yerine kural tabanlı bir motor aracılığıyla doğrudan işletmecinin telefonuna yapay zeka satın alma tavsiyesi olarak iletilmektedir."
add_body_paragraph(c2a, text_2a, is_first=True)

# 2.B. YÖNTEM
add_header_row(unified_table, "2.B. YÖNTEM / METHOD [PUAN / SCORE : 20]")
c2b = add_content_row(unified_table)
text_2b_p1 = "Projede izlenen yöntem ve teknik süreçler birbirini tamamlayan üç ana aşamada yürütülmüştür. İlk aşamada sistem mimarisi ve veri tabanı tasarımı ele alınmıştır. Backend tarafında asenkron çalışma desteği ve yüksek işlem performansı sunan Python tabanlı FastAPI çatısı tercih edilmiştir. Veri katmanında masalar, siparişler, menü bileşenleri, ürün varyasyonları, reçeteler ve hammadde stoklarını ilişkilendiren ilişkisel veri şeması oluşturulmuştur. İstemciler ile sunucu arasındaki canlı sipariş ve mutfak durumu güncellemeleri REST API ve WebSocket bağlantıları ile kurgulanmıştır."
text_2b_p2 = "İkinci aşamada veri hazırlama ve model eğitimi gerçekleştirilmiştir. Sistemin talep tahmin doğruluğunu test etmek amacıyla bir kafenin altı aylık operasyonunu yansıtan 37 binden fazla gerçekçi satış verisi üretilmiştir. Bu verilere Open-Meteo servisi üzerinden geçmiş hava durumu verileri eklenmiştir. Veri seti üzerinde tarih, haftanın günü, hafta sonu durumu, saat dilimi ve hava değişkenleri kullanılmıştır. Talep tahmini için hızlı çalışması ve düşük kaynak tüketimi nedeniyle LightGBM algoritması seçilmiş, ürün bazlı modeller eğitilmiş ve elde edilen modeller sisteme entegre edilmiştir."
text_2b_p3 = "Üçüncü aşamada reçete motoru ve karar destek mantığı kurulmuştur. Satış anında tetiklenen reçete motoru, satılan her ürünün içeriğindeki hammadde miktarlarını hesaplayarak depodan düşmektedir. Karar motoru ise LightGBM modelinden gelen saatlik satış tahmini ile reçete tüketim formüllerini birleştirerek ilgili hammaddelerin tahmini tükenme hızını belirlemektedir. Mevcut stok seviyesi kritik eşiğin altına indiğinde veya yaklaşan talep mevcut stoku aşacağı hesaplandığında, Telegram Bot servisi işletmeciye bildirim iletilmektedir."
add_body_paragraph(c2b, text_2b_p1, is_first=True)
add_body_paragraph(c2b, text_2b_p2, is_first=False)
add_body_paragraph(c2b, text_2b_p3, is_first=False)

# Add Flowchart Image under Yöntem
p_flow_img = c2b.add_paragraph()
p_flow_img.paragraph_format.space_before = Pt(8)
p_flow_img.paragraph_format.space_after = Pt(2)
p_flow_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_flow_img = p_flow_img.add_run()
run_flow_img.add_picture(flowchart_img, width=Inches(5.6))

# Add Caption for Flowchart
p_flow_cap = c2b.add_paragraph()
p_flow_cap.paragraph_format.space_before = Pt(2)
p_flow_cap.paragraph_format.space_after = Pt(4)
p_flow_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_flow_cap = p_flow_cap.add_run("Şekil 1. Projede izlenen yöntem ve sistem veri akış şeması.")
run_flow_cap.font.name = "Times New Roman"
run_flow_cap.font.size = Pt(10)
run_flow_cap.font.italic = True

# 2.C. ÇALIŞMA TAKVİMİ
add_header_row(unified_table, "2.C. ÇALIŞMA TAKVİMİ / cALENDER OF STUDY [PUAN / SCORE : 5]")
c2c = add_content_row(unified_table)

# Add Gantt Chart Image
p_gantt_img = c2c.paragraphs[0]
p_gantt_img.paragraph_format.space_before = Pt(4)
p_gantt_img.paragraph_format.space_after = Pt(2)
p_gantt_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_gantt_img = p_gantt_img.add_run()
run_gantt_img.add_picture(gantt_img, width=Inches(6.4))

# Caption: (Grafik-1)
p_gantt_cap = c2c.add_paragraph()
p_gantt_cap.paragraph_format.space_before = Pt(2)
p_gantt_cap.paragraph_format.space_after = Pt(6)
p_gantt_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_gantt_cap = p_gantt_cap.add_run("(Grafik-1)")
run_gantt_cap.font.name = "Times New Roman"
run_gantt_cap.font.size = Pt(10)

# Explanation text
add_body_paragraph(c2c, "Grafik-1'de yer alan unsurlar:", is_first=False, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
add_body_paragraph(c2c, "• Süreç ekseni: Ekim'den Mayıs'a 8 aylık periyot.", is_first=False, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
add_body_paragraph(c2c, "• Görev kategorileri: Hazırlık ve teslim (gri), Çekirdek motor ve modeller (koyu mavi), Genişletilmiş arayüzler (açık mavi), Genişletilmiş koşullu Telegram entegrasyonu (çizgili desen).", is_first=False, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
add_body_paragraph(c2c, "• Karar kontrol noktaları: K1 (Ekim sonu), K2 (Ocak sonu), K3 (Mart sonu) ve K4 (Nisan sonu) düşey kesikli çizgileri ve alt kısımdaki baklava simgeleri.", is_first=False, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)

# 2.D. KAYNAKÇA
add_header_row(unified_table, "2.D. KaynakÇA / REFERENCES [PUAN / SCORE : 5]")
c2d = add_content_row(unified_table)
refs = [
    "Aksoy, M. ve Çelik, E. (2024). Restoran işletmelerinde mobil tabanlı karar destek sistemlerinin operasyonel verimliliğe etkisi. Bilişim Teknolojileri Dergisi, 17(2), 145-156.",
    "Chen, H., Liu, Y. ve Wang, K. (2024). Impact of meteorological variables on short-term food service demand forecasting. International Journal of Hospitality Management, 118, 103-112.",
    "Kaya, S. ve Yıldız, A. (2023). Yeme-içme sektöründe bulut tabanlı satış noktası sistemlerinin envanter yönetimi üzerine etkileri. Gastronomi ve İşletme Araştırmaları Dergisi, 4(1), 35-48.",
    "Ke, G., Meng, Q., Finley, T., Wang, T., Chen, W., Ma, W., Ye, Q. ve Liu, T. Y. (2017). LightGBM: A highly efficient gradient boosting decision tree. Advances in Neural Information Processing Systems, 30, 3146-3154.",
    "Türker, G., Doğan, O. ve Şahin, B. (2026). Küçük ve orta ölçekli kafe işletmelerinde reçete bazlı stok kontrolü ve gıda israfını önleme yaklaşımları. Gıda Mühendisliği ve İşletme Yönetimi Dergisi, 8(1), 22-34.",
    "Wang, Z. ve Zhang, L. (2025). Machine learning approaches for daily demand forecasting in boutique cafes using weather and calendar features. Journal of Foodservice Business Research, 28(3), 289-307."
]
for idx, ref_text in enumerate(refs):
    add_body_paragraph(c2d, ref_text, is_first=(idx == 0), space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)

# 3.A. PROJE SUNUMU
add_header_row(unified_table, "3.A. PROJE SUNUMU / PROJECT PRESENTATION [PUAN / SCORE : 10]")
c3a = add_content_row(unified_table)
add_body_paragraph(c3a, "Bu başlık altına herhangi bir içerik yazmayınız. Raporunuzun değerlendirme sürecinde YAPILAN SÖZLÜ SUNUM (DEMO) PERFORMANSINA GÖRE değelendirilecektir.", is_first=True)

# 3.B. PROJE RAPOR DÜZENİ
add_header_row(unified_table, "3.B. PROJE RAPOR DÜZENİ / PROJECT REPORT WRITING RULES [PUAN / SCORE : 10]")
c3b = add_content_row(unified_table)
add_body_paragraph(c3b, "Bu başlık altına herhangi bir içerik yazmayınız. Raporunuzun değerlendirme sürecinde yazım kurallarına UYULMASI, rapor düzeni ve şablonA OLAN uyumu, proje fikrinizİN RAPOR İÇERİĞİNE NET BİR ŞEKİLDE AKTARILMASI, PUANLAMA KRİTERİ ALTINDA YER ALAN AÇIKLAMALARIN YETERLİ VE DOĞRU OLMASI GİBİ KRİTERLERE GÖRE Bu bölüm değelendirilecektir.", is_first=True)

# 4. EKLER
add_header_row(unified_table, "4. EKLER / aPPENDIX")
c4 = add_content_row(unified_table)
p_empty = c4.paragraphs[0]
p_empty.paragraph_format.space_before = Pt(0)
p_empty.paragraph_format.space_after = Pt(0)
r_empty = p_empty.add_run("")
r_empty.font.name = "Times New Roman"

doc.save(dst_path)
print(f"Duzenlenmis belge basariyla kaydedildi: {dst_path}")
