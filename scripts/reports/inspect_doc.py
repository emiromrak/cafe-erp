from _paths import IMAGE_DIR, OUTPUT_PATH, TEMPLATE_PATH

import docx

doc = docx.Document(TEMPLATE_PATH)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")

for i, p in enumerate(doc.paragraphs):
    text_preview = p.text.strip().replace('\n', ' ')[:60]
    has_image = len(p._element.xpath('.//a:blip')) > 0 or len(p._element.xpath('.//w:drawing')) > 0
    print(f"P[{i}]: text='{text_preview}' len={len(p.text)} runs={len(p.runs)} img={has_image}")

print("\n--- TABLES ---")
for t_idx, t in enumerate(doc.tables):
    print(f"\nTable {t_idx}: {len(t.rows)} rows x {len(t.columns)} cols")
    for r_idx, row in enumerate(t.rows):
        row_text = [cell.text.strip().replace('\n', ' ')[:30] for cell in row.cells]
        # show unique cell texts (since merged cells repeat)
        unique_cells = []
        for c in row_text:
            if not unique_cells or c != unique_cells[-1]:
                unique_cells.append(c)
        print(f"  Row {r_idx}: {unique_cells}")
